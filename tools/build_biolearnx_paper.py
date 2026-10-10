from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "paper"
ASSETS = OUT / "assets"
DOCX_PATH = OUT / "BioLearnX_ICAC_2026_Manuscript_Draft.docx"
FIG_PATH = ASSETS / "biolearnx_architecture.png"

FONT = "Times New Roman"
INK = RGBColor(0, 0, 0)


def set_run_font(run, size: float, *, bold=False, italic=False, small_caps=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = INK
    run._element.get_or_add_rPr()
    rfonts = run._element.rPr.rFonts
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    if small_caps:
        sc = OxmlElement("w:smallCaps")
        run._element.rPr.append(sc)
    return run


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")


def shade_cell(cell, fill="E7E7E7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_columns(section, count=2, space_twips=288):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols_el = cols[0]
    else:
        cols_el = OxmlElement("w:cols")
        sect_pr.append(cols_el)
    cols_el.set(qn("w:num"), str(count))
    cols_el.set(qn("w:space"), str(space_twips))
    cols_el.set(qn("w:equalWidth"), "1")


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_runs(paragraph, parts: Iterable[tuple[str, dict]], size=10):
    for text, opts in parts:
        set_run_font(paragraph.add_run(text), size, **opts)


def body_paragraph(doc, text: str, *, first_line=True, size=10, after=0.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.first_line_indent = Cm(0.35) if first_line else Cm(0)
    p.paragraph_format.widow_control = True
    set_run_font(p.add_run(text), size)
    return p


def heading1(doc, numeral: str, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    label = f"{numeral}. {text.upper()}" if numeral else text.upper()
    set_run_font(p.add_run(label), 10, small_caps=True)
    return p


def heading2(doc, letter: str, text: str):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(f"{letter}. {text}"), 10, italic=True)
    return p


def caption(doc, text: str, *, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3 if above else 1)
    p.paragraph_format.space_after = Pt(1 if above else 3)
    p.paragraph_format.keep_with_next = above
    set_run_font(p.add_run(text), 8)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        shade_cell(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), 7.5, bold=True)
        if widths:
            cell.width = Inches(widths[i])
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            set_run_font(p.add_run(value), 7.25)
            if widths:
                cell.width = Inches(widths[i])
    return table


def find_font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def make_architecture_figure(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    w, h = 2100, 760
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_font = find_font(42, True)
    box_font = find_font(34, True)
    small_font = find_font(28)
    tiny_font = find_font(25)
    line = (35, 35, 35)
    fill = (242, 242, 242)
    accent = (220, 227, 234)

    d.text((w // 2, 24), "Curriculum-Grounded Personalization Loop", font=title_font, fill=line, anchor="ma")
    boxes = [
        (60, 120, 510, 390, "1  Visual Learning", ["FLUX generation", "dual evaluation", "2-D/3-D interaction"]),
        (560, 120, 1010, 390, "2  Voice Learning", ["multilingual STT", "hybrid retrieval", "localized TTS"]),
        (1060, 120, 1510, 390, "3  Adaptive Assessment", ["grounded questions", "progressive hints", "weak-topic detection"]),
        (1560, 120, 2010, 390, "4  Note Digitization", ["Sinhala TrOCR", "image enhancement", "Tamil/English MT"]),
    ]
    for idx, (x1, y1, x2, y2, title, lines) in enumerate(boxes):
        d.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=accent if idx in (0, 2) else fill, outline=line, width=4)
        d.text(((x1 + x2) // 2, y1 + 42), title, font=box_font, fill=line, anchor="ma")
        yy = y1 + 118
        for item in lines:
            d.text((x1 + 38, yy), f"- {item}", font=small_font, fill=line)
            yy += 53

    for x in (535, 1035, 1535):
        d.line((x - 18, 255, x + 18, 255), fill=line, width=6)
        d.polygon([(x + 18, 255), (x + 2, 243), (x + 2, 267)], fill=line)

    d.rounded_rectangle((355, 500, 1745, 690), radius=24, fill=(250, 250, 250), outline=line, width=4)
    d.text((1050, 530), "Shared learner and curriculum context", font=box_font, fill=line, anchor="ma")
    d.text((1050, 595), "source material  |  learner history  |  feedback  |  generated artifacts", font=small_font, fill=line, anchor="ma")
    d.text((1050, 648), "Component stores exist; a production cross-component gateway remains to be validated.", font=tiny_font, fill=line, anchor="ma")
    for x in (285, 785, 1285, 1785):
        d.line((x, 390, x, 500), fill=line, width=4)
        d.polygon([(x, 500), (x - 12, 481), (x + 12, 481)], fill=line)
    img.save(path, dpi=(300, 300), optimize=True)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    normal._element.rPr.append(lang)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def build_document():
    OUT.mkdir(parents=True, exist_ok=True)
    make_architecture_figure(FIG_PATH)
    doc = Document()
    doc.core_properties.title = "BioLearnX: Curriculum-Grounded Multimodal Agents for Personalized Biology Learning"
    doc.core_properties.subject = "ICAC 2026 conference manuscript draft"
    doc.core_properties.author = "Kojithan P. Y.; Baskaran V.; Nishara T.; Sarmitha S."
    doc.core_properties.keywords = "multimodal learning, retrieval-augmented generation, adaptive assessment, multilingual speech, Sinhala OCR"
    configure_styles(doc)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(4.3)
    sec.left_margin = Cm(1.43)
    sec.right_margin = Cm(1.43)
    sec.header_distance = Cm(0.5)
    sec.footer_distance = Cm(0.5)
    set_columns(sec, 1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("BioLearnX: Curriculum-Grounded Multimodal Agents for Personalized Biology Learning"), 24)

    authors = [
        ("Kojithan P. Y.", "it22264220@my.sliit.lk"),
        ("Baskaran V.", "it22172600@my.sliit.lk"),
        ("Nishara T.", "it22223876@my.sliit.lk"),
        ("Sarmitha S.", "it22637482@my.sliit.lk"),
    ]
    author_table = doc.add_table(rows=2, cols=2)
    author_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    author_table.autofit = False
    tbl_pr = author_table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)
    for i, (name, email) in enumerate(authors):
        cell = author_table.cell(i // 2, i % 2)
        cell.width = Inches(3.45)
        set_cell_margins(cell, top=20, start=40, bottom=20, end=40)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for line_no, line in enumerate((name, "Faculty of Computing", "Sri Lanka Institute of Information Technology", "Malabe, Sri Lanka", email)):
            set_run_font(p.add_run(line), 11 if line_no == 0 else 10.2, bold=(line_no == 0))
            if line_no < 4:
                p.add_run().add_break(WD_BREAK.LINE)

    abstract = doc.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.space_before = Pt(5)
    abstract.paragraph_format.space_after = Pt(3)
    abstract.paragraph_format.left_indent = Cm(0.7)
    abstract.paragraph_format.right_indent = Cm(0.7)
    abstract.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run_font(abstract.add_run("Abstract—"), 9, bold=True, italic=True)
    set_run_font(abstract.add_run(
        "Biology learning requires learners to coordinate visual structure, verbal explanation, assessment feedback, and handwritten study material. Existing tools commonly treat these activities as disconnected services, limiting the reuse of curriculum evidence and learner history. This paper presents BioLearnX, a multimodal framework that combines four independently developed components: reflective generation of interactive biological visuals, multilingual voice-based retrieval and explanation, curriculum-grounded adaptive assessment, and Sinhala handwritten-note digitization with Tamil and English translation. The components use retrieval-augmented generation and explicit quality gates to constrain generated content, while learner interactions form a proposed personalization loop. A repository audit found 236 Python source files across the four components to be syntactically parseable, and the visualization client passed static type checking. However, the supplied artifacts do not include the raw benchmark outputs or an end-to-end learner study needed to support comparative accuracy, latency, or learning-effectiveness claims. We therefore report the implemented architecture, a contribution-to-evaluation protocol, and the evidence required for a valid conference submission without manufacturing empirical results."
    ), 9, italic=True)

    keywords = doc.add_paragraph()
    keywords.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    keywords.paragraph_format.left_indent = Cm(0.7)
    keywords.paragraph_format.right_indent = Cm(0.7)
    keywords.paragraph_format.space_after = Pt(4)
    set_run_font(keywords.add_run("Keywords—"), 9, bold=True, italic=True)
    set_run_font(keywords.add_run("adaptive assessment, biology education, multimodal learning, multilingual speech, optical character recognition, retrieval-augmented generation"), 9, italic=True)

    fig_p = doc.add_paragraph()
    fig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_p.paragraph_format.space_before = Pt(2)
    fig_p.paragraph_format.space_after = Pt(0)
    fig_p.paragraph_format.keep_with_next = True
    figure_shape = fig_p.add_run().add_picture(str(FIG_PATH), width=Cm(17.8))
    figure_shape._inline.docPr.set("descr", "BioLearnX architecture linking visual learning, voice learning, adaptive assessment, and note digitization to shared learner and curriculum context.")
    caption(doc, "Fig. 1. BioLearnX architecture and the proposed curriculum-grounded personalization loop.")

    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    body_sec.page_width = Cm(21.0)
    body_sec.page_height = Cm(29.7)
    body_sec.top_margin = Cm(1.9)
    body_sec.bottom_margin = Cm(4.3)
    body_sec.left_margin = Cm(1.43)
    body_sec.right_margin = Cm(1.43)
    set_columns(body_sec, 2, 288)

    heading1(doc, "I", "Introduction")
    body_paragraph(doc, "Biology education combines spatial anatomy, causal processes, terminology, and cumulative recall. Multimedia explanations can support learning when verbal and visual representations are coordinated [1], while retrieval practice can improve long-term retention [2]. Adaptive testing can further select tasks that match a learner's evolving ability [3]. In practice, however, visual generators, speech tutors, quiz systems, and note digitizers are often isolated. Their outputs are not consistently grounded in the same learning material, and feedback collected in one modality rarely informs another.")
    body_paragraph(doc, "BioLearnX addresses this fragmentation through four coordinated components developed as one research project: (1) interactive biological visualization, (2) multilingual audio communication, (3) adaptive quiz generation and recommendation, and (4) Sinhala handwritten-note extraction and translation. Fig. 1 shows the intended loop. Each component is usable independently, but the research contribution is their shared curriculum-grounded design and the proposed circulation of source material, learner performance, and feedback.")
    body_paragraph(doc, "This paper asks four research questions. RQ1 examines whether reflective generation improves visual and anatomical quality over a linear generator. RQ2 examines whether source grounding and adaptive rules improve question validity and recommendation relevance. RQ3 examines the accuracy-latency trade-offs of multilingual speech and Sinhala OCR pipelines. RQ4 examines whether the integrated loop improves learning, usability, and accessibility relative to isolated tools. The contributions are: an architecture unifying four learning modalities; explicit grounding and quality gates within each pipeline; an evaluation protocol linking each contribution to measurable evidence; and a transparent audit of the evidence currently available.")

    heading1(doc, "II", "Related Work")
    heading2(doc, "A", "Grounded and adaptive learning")
    body_paragraph(doc, "Retrieval-augmented generation (RAG) combines parametric generation with evidence retrieved from an external collection [4]. BioLearnX uses this pattern to ground explanations and questions in learner-provided or curriculum material. Sentence-BERT supports efficient semantic retrieval [5], and BGE-M3 extends retrieval across languages and retrieval functions [6]. In the audio component, dense and lexical ranks are combined using reciprocal rank fusion (RRF), a robust method for merging ranked lists [7]. Unlike a single RAG chatbot, BioLearnX applies grounding at several pedagogical decision points: explanation, question construction, feedback, and recommendation.")
    heading2(doc, "B", "Multimodal generation and reflection")
    body_paragraph(doc, "Vision-language alignment models such as CLIP provide a machine-checkable relation between images and text [8]. Segment Anything 2 enables promptable segmentation across images and video [9], Qwen2.5-VL supports vision-language reasoning [10], and Hunyuan3D 2.0 provides a route from generated imagery to textured three-dimensional assets [11]. BioLearnX combines these capabilities with a dual visual-pedagogical critic and a bounded reflection loop. This follows the broader idea of agents improving subsequent actions from explicit feedback [12], but specializes the feedback to biological structure and educational usefulness.")
    heading2(doc, "C", "Low-resource speech and document understanding")
    body_paragraph(doc, "Whisper established a broadly multilingual speech-recognition baseline [13], while NLLB addresses machine translation over 200 languages [14]. For handwritten notes, TrOCR formulates recognition as an image-to-text Transformer task [15], and Swin2SR targets image restoration and super-resolution [16]. BioLearnX routes speech and OCR models by language and task rather than assuming one model is uniformly reliable for English, Tamil, and Sinhala. Its assessment component maps learning objectives to Bloom-style cognitive levels [17] and uses learner attempts and scores to adjust difficulty.")

    heading1(doc, "III", "System Design and Methodology")
    heading2(doc, "A", "Integrated architecture")
    body_paragraph(doc, "The system is implemented as four service-oriented applications. Each accepts a learning artifact or interaction, applies modality-specific processing, and returns a learner-facing artifact together with metadata needed for later feedback. The shared design principle is evidence before generation: curriculum or learner material is ingested, indexed, and retrieved before a model produces an explanation, question, recommendation, or translated note. Current repositories retain component-level state; a production cross-component identity, consent, and event gateway is part of the integration protocol rather than a verified deployment feature.")
    caption(doc, "TABLE I\nIMPLEMENTED COMPONENTS AND EVIDENCE BOUNDARIES", above=True)
    add_table(doc,
        ["Component", "Implemented pipeline", "Quality or grounding control"],
        [
            ["Visual", "Prompt enhancement; FLUX image generation; optional segmentation, labels, and 3-D conversion", "Dual visual/pedagogical scores; anatomy hard-failure checks; at most two reflective retries"],
            ["Audio", "Language-routed STT; document RAG; localization; language-routed TTS", "User-scoped retrieval; dense and BM25 fusion; reranking; top-five context"],
            ["Assessment", "Document ingestion; grounded quiz generation; evaluation; hints; recommendation", "Grounding threshold 0.55; fail-closed generation; adaptive attempt/score rules"],
            ["Notes", "Image validation; Sinhala TrOCR; Swin2SR comparison asset; Tamil/English translation", "OCR uses original pixels in the active route; enhancement is returned separately; context improvement disabled"],
        ], [0.75, 1.25, 1.25])

    heading2(doc, "B", "Interactive biological visualization")
    body_paragraph(doc, "A prompt agent expands a learner request with biological and pedagogical constraints before a FLUX.1-dev generator creates the initial image. Visual and pedagogical evaluators score the artifact, while anatomy-specific catalogs check structure, relation, orientation, canonical labels, overlap, crossing, and hard failures for supported organs. An artifact must achieve both visual and pedagogical scores of at least 7/10 and avoid anatomy hard failures; otherwise a reflection agent rewrites the prompt and regeneration is attempted, with two retries permitted. Accepted artifacts can be enriched through SAM2/Qwen2.5-VL interaction generation, SVG labels, or Hunyuan3D conversion. Retrieved curriculum context uses MiniLM embeddings, PostgreSQL vector and full-text search, and RRF. Persistent feedback is gated by repeated evidence, held-out evaluation, and rollback criteria.")

    heading2(doc, "C", "Multilingual audio communication")
    body_paragraph(doc, "The audio workflow is a linear state graph: speech recognition, prompt enhancement, retrieval, localization, and speech synthesis. Sinhala uses a dedicated Whisper-small model; Tamil uses Qwen3 ASR with Whisper Large V3 fallback; other routes use Whisper Large V3. Uploaded PDF, presentation, word-processing, spreadsheet, text, or Markdown content is chunked into 500-character windows with 50-character overlap. Dense BGE-M3 and BM25 retrieval run in parallel, their ranks are fused with RRF using k=60, results are reranked, and the five highest-ranked passages form the answer context. The grounded answer is produced in English, localized to the target language, and synthesized using Kokoro for English, IndicF5 for Tamil, or VITS for Sinhala. These routing statements describe code paths, not measured accuracy.")

    heading2(doc, "D", "Adaptive assessment and recommendation")
    body_paragraph(doc, "Assessment documents are chunked into 800-character windows with 150-character overlap and embedded using all-MiniLM-L6-v2. Generated multiple-choice, structured, and essay questions are accepted only when grounding reaches 0.55; a failed grounding gate does not silently release the question. For multiple-choice and fill-in tasks, at most two terminal attempts select hard difficulty, three select medium, and four or more select easy. Structured or essay scores of at least 0.70 select hard, scores from 0.40 to 0.69 select medium, and lower scores select easy. Three progressive hints are available before reveal or progression. A topic is marked weak when average performance is below 0.60 or average terminal attempts exceed one, and recommendations are tied to retrieved source chunks.")

    heading2(doc, "E", "Sinhala note digitization")
    body_paragraph(doc, "The active note-processing endpoint validates JPEG, PNG, WebP, or BMP input. Swin2SR creates a fourfold enhanced image for comparison and download, but fine-tuned TrOCR reads the original pixels because enhancement can change handwriting strokes. OCR text is cleaned for Sinhala noise and translated sequentially to Tamil and English with NLLB-200. The active route returns a whole-page result and explicitly reports that context improvement is disabled. This distinction is important: older design notes describe line-level processing and language-model correction, but they are not treated as implemented behavior in this study.")

    heading1(doc, "IV", "Evaluation Protocol")
    heading2(doc, "A", "Experimental questions, baselines, and metrics")
    body_paragraph(doc, "The protocol separates component quality from end-to-end learning impact. For RQ1, the repository defines a 100-prompt visual suite, five random seeds, and six configurations: a linear generator; dual critics; critics plus reflection; critics plus persistent memory; critics plus skill instructions; and the full configuration. Report visual and pedagogical scores, anatomy structure/relation/orientation metrics, hard-failure rate, acceptance rate, retry count, and latency. Pair identical prompts and seeds, report means with standard deviations and 95% confidence intervals, and test the full system against the strongest ablation.")
    body_paragraph(doc, "For RQ2, construct an expert-reviewed corpus spanning biology topics and Bloom levels. Compare dense-only, lexical-only, fused, and reranked retrieval using recall at k and mean reciprocal rank. Evaluate question grounding, factual validity, distractor plausibility, Bloom alignment, and recommendation relevance with at least two independent biology reviewers; report agreement and adjudication. Test the adaptive policy against fixed-difficulty and random-question controls, using pre-test/post-test gain, completion time, hint use, and delayed retention.")
    body_paragraph(doc, "For RQ3, use held-out, speaker-separated English, Tamil, and Sinhala audio with verified transcripts. Report word error rate, character error rate where appropriate, real-time factor, end-to-end latency, translation adequacy, and speech intelligibility. For OCR, the benchmark script expects 227 held-out images and compares base TrOCR, fine-tuned TrOCR, Swin2SR plus fine-tuned TrOCR, and the active application core. Report character error rate (CER), word error rate (WER), exact match, and warm-GPU latency. For RQ4, conduct a counterbalanced learner study comparing the integrated loop with isolated components; pre-register outcomes and obtain institutional ethics approval before collecting learner data.")
    caption(doc, "TABLE II\nCONTRIBUTION-TO-EVALUATION TRACEABILITY AND CURRENT STATUS", above=True)
    add_table(doc,
        ["RQ", "Primary comparison", "Required evidence", "Current status"],
        [
            ["RQ1", "Full visual agent vs. strongest ablation", "Paired prompt/seed outputs; critic and anatomy scores", "Protocol defined; outputs absent"],
            ["RQ2", "Grounded adaptive policy vs. fixed/random controls", "Expert labels; retrieval judgments; learning and retention outcomes", "Implementation present; study outputs absent"],
            ["RQ3", "Language/model routes and OCR ablations", "Verified transcripts; held-out OCR ground truth; raw latency logs", "Benchmark code present; auditable outputs absent"],
            ["RQ4", "Integrated loop vs. isolated components", "Ethics approval; learner allocation; pre/post and usability data", "Cross-component learner study not supplied"],
        ], [0.32, 0.93, 1.05, 0.96])

    heading2(doc, "B", "Reproducibility and statistical reporting")
    body_paragraph(doc, "All experiments should record model identifiers and revisions, prompts, decoding parameters, hardware, software versions, dataset licenses, random seeds, per-item predictions, timing boundaries, and exclusion rules. Performance values should be computed from raw per-item records, not screenshots. Confidence intervals and effect sizes should accompany significance tests, and multiple comparisons should be corrected where applicable. Human evaluations require blinded ordering, a written rubric, reviewer qualifications, and inter-rater agreement. Personally identifiable learner material must be minimized, access-controlled, and deleted according to the approved protocol.")

    heading1(doc, "V", "Results and Evidence Status")
    heading2(doc, "A", "Repository-level verification")
    body_paragraph(doc, "A static repository audit parsed 236 Python files without syntax failure: 72 in the visualization component, 73 in audio communication, 56 in adaptive assessment, and 35 in note digitization. Virtual environments, caches, and generated files were excluded. The visualization web client also passed its configured TypeScript static type check. These checks support the claim that the inspected source tree is structurally parseable; they do not establish runtime correctness, model quality, retrieval accuracy, latency, or educational benefit.")
    heading2(doc, "B", "Empirical evidence available for the research questions")
    body_paragraph(doc, "No raw experiment record found in the supplied repositories resolves RQ1-RQ4. The visual component includes evaluation configurations but not completed JSONL outputs. Audio contains a concurrency unit test with mocked delays, which cannot be reported as system latency. The assessment component contains grounding and adaptation logic but no expert-rated question set or learner outcomes. The OCR component contains an evaluation script and a summary graph that states relative improvements, but the held-out images, predictions, and per-sample metric outputs are absent; those values are therefore excluded from the findings. Table II identifies the evidence required before submission.")
    heading2(doc, "C", "Interpretation")
    body_paragraph(doc, "The available evidence supports an implementation contribution and a testable integrated design, not a comparative performance or learning-effectiveness claim. The strongest defensible conclusion is that four modality-specific pipelines implement a common evidence-grounded pattern and expose measurable quality gates. Acceptance-oriented claims such as improved anatomy accuracy, lower CER/WER, faster response, or better learning must remain contingent on the protocol in Section IV. Reporting this boundary prevents implementation completeness from being mistaken for empirical validation.")

    heading1(doc, "VI", "Discussion")
    heading2(doc, "A", "Why the integration is technically meaningful")
    body_paragraph(doc, "The four components address complementary failure modes. Visual reflection can reject anatomically implausible diagrams; retrieval can bind spoken explanations to source material; adaptive assessment can convert learner performance into weak-topic evidence; and OCR can admit Sinhala handwritten notes into the same learning workflow. Their combination is more than a menu of tools only if identity, curriculum provenance, consent, and learner events are shared reliably. The current component stores make this integration feasible, but the shared gateway and causal benefit of cross-modal personalization still require validation.")
    heading2(doc, "B", "Threats to validity")
    body_paragraph(doc, "Construct validity depends on whether automated critics and grounding scores represent pedagogical quality. Internal validity is threatened by model nondeterminism, prompt leakage between tuning and evaluation, and uncontrolled hardware timing. External validity is limited if prompts, speakers, handwriting, or learners do not represent Sri Lankan curricula and language varieties. Reviewer bias can affect question and translation ratings. Finally, component-specific gains may not transfer to learning outcomes; only a controlled learner study can establish that relationship.")
    heading2(doc, "C", "Ethical and deployment considerations")
    body_paragraph(doc, "Biology diagrams and explanations can be confidently wrong, speech and handwriting models can underperform for low-resource varieties, and learner profiles can expose sensitive educational data. Deployment should therefore display source provenance, preserve human correction, log model and retrieval versions, and provide accessible alternatives. Learner studies require informed consent, data minimization, secure retention, and an explicit procedure for withdrawing records. Automated recommendations should support, not replace, instructor judgment.")

    heading1(doc, "VII", "Conclusion")
    body_paragraph(doc, "BioLearnX integrates interactive biological visualization, multilingual voice learning, adaptive grounded assessment, and Sinhala note digitization within a proposed curriculum-grounded personalization loop. The implementation audit confirms a substantial, structurally parseable prototype and identifies explicit gates for image quality, retrieval, question grounding, and OCR routing. The manuscript does not claim unobserved performance: raw component benchmarks and an end-to-end learner study are required to answer the four research questions. Completing the paired ablations, multilingual accuracy tests, expert review, and controlled learner evaluation will convert the current systems contribution into an empirically supported ICAC submission.")

    heading1(doc, "", "Acknowledgment")
    body_paragraph(doc, "The authors thank Nuwan Kodagoda and Malithi Nawarathne of the Faculty of Computing, Sri Lanka Institute of Information Technology, for supervising the project.", first_line=False)

    heading1(doc, "", "References")
    refs = [
        "R. E. Mayer, Multimedia Learning, 2nd ed. Cambridge, U.K.: Cambridge Univ. Press, 2009.",
        "H. L. Roediger III and J. D. Karpicke, \"Test-enhanced learning: Taking memory tests improves long-term retention,\" Psychological Science, vol. 17, no. 3, pp. 249-255, 2006.",
        "W. J. van der Linden and C. A. W. Glas, Eds., Elements of Adaptive Testing. New York, NY, USA: Springer, 2010.",
        "P. Lewis et al., \"Retrieval-augmented generation for knowledge-intensive NLP tasks,\" in Adv. Neural Inf. Process. Syst., vol. 33, 2020, pp. 9459-9474.",
        "N. Reimers and I. Gurevych, \"Sentence-BERT: Sentence embeddings using Siamese BERT-networks,\" in Proc. EMNLP-IJCNLP, 2019, pp. 3982-3992.",
        "J. Chen et al., \"BGE M3-Embedding: Multi-lingual, multi-functionality, multi-granularity text embeddings through self-knowledge distillation,\" arXiv:2402.03216, 2024.",
        "G. V. Cormack, C. L. A. Clarke, and S. Buettcher, \"Reciprocal rank fusion outperforms Condorcet and individual rank learning methods,\" in Proc. 32nd ACM SIGIR Conf., 2009, pp. 758-759.",
        "A. Radford et al., \"Learning transferable visual models from natural language supervision,\" in Proc. ICML, vol. 139, 2021, pp. 8748-8763.",
        "N. Ravi et al., \"SAM 2: Segment anything in images and videos,\" arXiv:2408.00714, 2024.",
        "S. Bai et al., \"Qwen2.5-VL technical report,\" arXiv:2502.13923, 2025.",
        "Z. Zhao et al., \"Hunyuan3D 2.0: Scaling diffusion models for high resolution textured 3D assets generation,\" arXiv:2501.12202, 2025.",
        "N. Shinn et al., \"Reflexion: Language agents with verbal reinforcement learning,\" in Adv. Neural Inf. Process. Syst., vol. 36, 2023.",
        "A. Radford et al., \"Robust speech recognition via large-scale weak supervision,\" in Proc. ICML, vol. 202, 2023, pp. 28492-28518.",
        "M. R. Costa-jussa et al., \"No language left behind: Scaling human-centered machine translation,\" arXiv:2207.04672, 2022.",
        "M. Li et al., \"TrOCR: Transformer-based optical character recognition with pre-trained models,\" in Proc. AAAI Conf. Artif. Intell., vol. 37, no. 11, 2023, pp. 13094-13102.",
        "M. V. Conde, U.-J. Choi, M. Burchi, and R. Timofte, \"Swin2SR: SwinV2 transformer for compressed image super-resolution and restoration,\" in Proc. ECCV Workshops, 2022, pp. 669-687.",
        "L. W. Anderson and D. R. Krathwohl, Eds., A Taxonomy for Learning, Teaching, and Assessing: A Revision of Bloom's Taxonomy of Educational Objectives. New York, NY, USA: Longman, 2001.",
    ]
    for idx, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        set_run_font(p.add_run(f"[{idx}] {ref}"), 8)

    # Remove any page-number fields and keep the manuscript visually clean.
    for section in doc.sections:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
        for p in list(section.header.paragraphs) + list(section.footer.paragraphs):
            p.clear()

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
